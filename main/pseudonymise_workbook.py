"""Pseudonymise every sheet of the thesis workbook, in place.

The committed `dataset/final_dataset.xlsx` carries identifiers across all six of
its sheets, not only the 10k sheet used for training: 13,038 @handle occurrences
(7,597 distinct) plus 3,818 URLs. Publishing that pairs live accounts with a
mental-health label, so the workbook that is committed needs to stop carrying
real handles.

This rewrites every string cell that matches a handle or a profile URL, using
the same salt and map as `build_public_release.py`, so the workbook and the
published JSONL release stay identical in text.

Deliberately minimal: only identifiers change. Text, labels, sheet names and
cell layout are otherwise untouched, so the workbook remains the file the
reported metrics were computed from, apart from the handle substitution. Emoji
mojibake in this workbook is left alone for the same reason - repairing it is
available in build_public_release.py behind an explicit flag, and is a known
limitation rather than a silent change here.

Usage:
    python main/pseudonymise_workbook.py --check     # measure only
    python main/pseudonymise_workbook.py --write     # rewrite in place
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

from build_public_release import (
    HANDLE_RE,
    PROFILE_URL_RE,
    VERIFY_RE,
    load_salt,
    pseudonymise,
)

WORKBOOK = Path("dataset/final_dataset.xlsx")
SALT_FILE = Path("private/pseudonym_salt.txt")
MAP_FILE = Path("private/handle_map.jsonl")


def load_map(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    if path.exists():
        for line in path.read_text().splitlines():
            if line.strip():
                rec = json.loads(line)
                mapping[rec["handle"]] = rec["pseudonym"]
    return mapping


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--write", action="store_true", help="rewrite the workbook in place")
    ap.add_argument("--check", action="store_true", help="measure only (default)")
    ap.add_argument("--workbook", type=Path, default=WORKBOOK)
    ap.add_argument("--salt-file", type=Path, default=SALT_FILE)
    ap.add_argument("--map-file", type=Path, default=MAP_FILE)
    args = ap.parse_args()

    salt = load_salt(args.salt_file)
    mapping = load_map(args.map_file)
    print(f"loaded {len(mapping)} existing pseudonyms")

    wb = openpyxl.load_workbook(args.workbook)
    before = {"handles": 0, "urls": 0}
    after = {"handles": 0, "urls": 0}
    cells_touched = 0
    residual_examples: list[tuple[str, str, str, str]] = []

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or "@" not in v and "http" not in v:
                    continue
                before["handles"] += len(HANDLE_RE.findall(v))
                before["urls"] += len(PROFILE_URL_RE.findall(v))
                new = pseudonymise(v, salt, mapping)
                if new != v:
                    # VERIFY_RE, not HANDLE_RE: the replacements are themselves
                    # @-shaped, so HANDLE_RE would count our own output.
                    after["handles"] += len(VERIFY_RE.findall(new))
                    after["urls"] += len(PROFILE_URL_RE.findall(new))
                    for m in VERIFY_RE.finditer(new):
                        if len(residual_examples) < 5:
                            residual_examples.append((
                                ws.title,
                                m.group(0)[:60],
                                repr(v[max(0, m.start() - 30):m.end() + 30]),
                                repr(new[max(0, m.start() - 30):m.end() + 30]),
                            ))
                    cells_touched += 1
                    if args.write:
                        cell.value = new

    print(f"sheets: {[ws.title for ws in wb.worksheets]}")
    print(f"handle occurrences before={before['handles']} urls_before={before['urls']}")
    print(f"cells rewritten: {cells_touched}")
    print(f"residual handles in rewritten cells: {after['handles']}")
    for sheet, tok, orig, newctx in residual_examples:
        print(f"  residual: sheet={sheet!r} token={tok!r}")
        print(f"     original: {orig}")
        print(f"     rewritten: {newctx}")

    if not args.write:
        print("\n--check only; nothing written. Re-run with --write to apply.")
        return 0

    if after["handles"] != 0:
        print("refusing to write: residual handles remain", file=sys.stderr)
        return 1

    wb.save(args.workbook)

    # Persist the map. Without this the pseudonyms created for handles that only
    # appear outside the 10k sheet exist nowhere, and those rows are not
    # reversible by the author.
    args.map_file.parent.mkdir(parents=True, exist_ok=True)
    with args.map_file.open("w") as fh:
        for handle, pseudo in sorted(mapping.items()):
            fh.write(json.dumps({"handle": handle, "pseudonym": pseudo}) + "\n")
    args.map_file.chmod(0o600)
    print(f"wrote {args.map_file} ({len(mapping)} pseudonyms)")
    print(f"wrote {args.workbook}")

    # verify by reloading from disk
    wb2 = openpyxl.load_workbook(args.workbook)
    residual = 0
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    residual += len(VERIFY_RE.findall(cell.value))
    print(f"VERIFICATION: residual handles after reload = {residual}")
    return 0 if residual == 0 else 1


if __name__ == "__main__":
    sys.exit(main())